from __future__ import annotations

import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
HERE = Path(__file__).resolve().parent
DATA = json.loads((HERE / "deck.json").read_text(encoding="utf-8"))
OUT = HERE / "low-alt-deck.pptx"
W, H = 13.333, 7.5
BG = RGBColor(8, 10, 12)
PANEL = RGBColor(17, 22, 22)
INK = RGBColor(240, 240, 233)
MUTED = RGBColor(180, 185, 177)
LINE = RGBColor(57, 64, 61)
ACID = RGBColor(216, 250, 85)
CYAN = RGBColor(141, 246, 238)

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]


def rect(slide, x, y, w, h, color, transparency=0, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background() if line is None else None
    if line is not None:
        shape.line.color.rgb = line
        shape.line.width = Pt(.8)
    return shape


def add_text(slide, text, x, y, w, h, size, color=INK, bold=False, align=PP_ALIGN.RIGHT, font="Arial", mono=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(.04)
    tf.margin_right = Inches(.04)
    tf.margin_top = Inches(.025)
    tf.margin_bottom = Inches(.025)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        p.line_spacing = 1.12
        p._p.get_or_add_pPr().set("rtl", "1" if align == PP_ALIGN.RIGHT else "0")
        run = p.add_run()
        run.text = line
        run.font.name = "Menlo" if mono else font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run._r.get_or_add_rPr().set("lang", "he-IL")
        run._r.get_or_add_rPr().set("rtl", "1" if align == PP_ALIGN.RIGHT else "0")
    return shape


def common(slide, index, eyebrow):
    rect(slide, 0, 0, W, H, BG)
    add_text(slide, "03   LOW / ALT", .58, .22, 2.4, .36, 10, ACID, True, PP_ALIGN.LEFT, mono=True)
    add_text(slide, eyebrow, 9.1, .27, 3.55, .3, 10, ACID, False, mono=True)
    rect(slide, .6, .72, W - 1.2, .012, LINE)
    add_text(slide, f"{index:02d} / {len(DATA['slides']):02d}", .6, 7.12, 1.4, .2, 9, MUTED, False, PP_ALIGN.LEFT, mono=True)


def add_crop(slide, path, x, y, w, h):
    p = ROOT / "presentation" / path
    iw, ih = Image.open(p).size
    pic = slide.shapes.add_picture(str(p), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    image_ratio, box_ratio = iw / ih, w / h
    if image_ratio > box_ratio:
        crop = (1 - box_ratio / image_ratio) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    else:
        crop = (1 - image_ratio / box_ratio) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop
    return pic


def title(slide, text, x=7.45, y=1.45, w=5.1, h=1.25, size=34):
    return add_text(slide, text, x, y, w, h, size, INK, True)


for index, item in enumerate(DATA["slides"], 1):
    slide = prs.slides.add_slide(blank)
    kind = item["type"]
    if kind == "cover":
        add_crop(slide, "assets/cover-art.png", 0, 0, W, H)
        add_text(slide, "03   LOW / ALT", .65, .38, 3, .4, 12, ACID, True, PP_ALIGN.LEFT, mono=True)
        add_text(slide, "LOW / ALT", .65, 3.05, 9.2, 1.15, 62, INK, True, PP_ALIGN.LEFT, mono=True)
        add_text(slide, item["subtitle"], .7, 4.35, 8.1, .55, 27, ACID, True)
        add_text(slide, item["body"], .7, 5.07, 7.2, .75, 16, MUTED)
        add_text(slide, item["eyebrow"], .7, 6.62, 7.2, .25, 10, MUTED, False, PP_ALIGN.LEFT, mono=True)
        continue

    common(slide, index, item["eyebrow"])
    if kind in ("split", "game"):
        add_crop(slide, item["image"], .63, 1.32, 6.65, 5.25)
        title(slide, item["title"], 7.65, 1.55, 4.95, 1.15, 33)
        y = 2.95
        if item.get("subtitle"):
            add_text(slide, item["subtitle"], 7.67, y, 4.9, .55, 19, ACID, True)
            y += .73
        add_text(slide, item["body"], 7.67, y, 4.85, 1.65, 18, MUTED)
        if item.get("controls"):
            rect(slide, 7.7, 5.56, 4.78, .53, PANEL, 0, LINE)
            add_text(slide, item["controls"], 7.86, 5.61, 4.45, .4, 12, CYAN, True, mono=True)
        if item.get("tags"):
            add_text(slide, "   ·   ".join(item["tags"]), 7.7, 5.7, 4.8, .65, 11, ACID)
        if item.get("note"):
            add_text(slide, item["note"], 7.7, 6.18, 4.8, .34, 10, MUTED, False, mono=True)
    elif kind == "timeline":
        title(slide, item["title"], 7.55, 1.3, 5.0, 1.15, 36)
        add_text(slide, item["body"], 7.65, 2.55, 4.8, .55, 16, MUTED)
        for i, row in enumerate(item["items"]):
            x = .63 + i * 3.07
            rect(slide, x, 3.58, 2.86, 2.53, PANEL, 0, LINE)
            add_text(slide, row[0], x + .18, 3.77, 2.45, .34, 12, ACID, True, PP_ALIGN.LEFT, mono=True)
            add_text(slide, row[1], x + .18, 4.33, 2.45, .45, 20, INK, True)
            add_text(slide, row[2], x + .18, 4.94, 2.45, .9, 14, MUTED)
    elif kind == "metrics":
        title(slide, item["title"], 7.55, 1.34, 5.0, .95, 36)
        for i, row in enumerate(item["items"]):
            x = .65 + i * 3.08
            rect(slide, x, 2.72, 2.84, 1.8, PANEL, 0, LINE)
            add_text(slide, row[0], x + .15, 2.9, 2.5, .78, 36, ACID, True, PP_ALIGN.LEFT, mono=True)
            add_text(slide, row[1], x + .15, 3.76, 2.5, .45, 13, MUTED)
        add_text(slide, item["body"], 1.0, 5.18, 11.2, .74, 17, INK)
    elif kind == "pack":
        title(slide, item["title"], 7.5, 1.3, 5.0, 1.2, 36)
        for i, row in enumerate(item["items"]):
            x = .67 + i * 4.14
            rect(slide, x, 3.0, 3.85, 1.8, PANEL, 0, LINE)
            add_text(slide, row[0], x + .2, 3.18, 3.43, .47, 13, ACID, True, PP_ALIGN.LEFT, mono=True)
            add_text(slide, row[1], x + .2, 3.8, 3.4, .66, 15, MUTED)
        add_text(slide, item["note"], 1.0, 5.55, 11.1, .7, 16, INK)
    elif kind == "checklist":
        title(slide, item["title"], 7.6, 1.25, 4.9, 1.1, 36)
        for i, text in enumerate(item["items"]):
            col, row = i % 2, i // 2
            x = .68 + col * 6.1
            y = 2.65 + row * 1.02
            rect(slide, x, y, 5.76, .77, PANEL, 0, LINE)
            add_text(slide, "✓", x + 5.25, y + .1, .3, .52, 16, ACID, True, PP_ALIGN.CENTER, mono=True)
            add_text(slide, text, x + .2, y + .08, 4.85, .57, 13, INK)
    elif kind == "release":
        title(slide, item["title"], 7.5, 1.4, 5.0, 1.2, 39)
        add_text(slide, item["body"], 7.65, 2.9, 4.8, 1.2, 19, MUTED)
        for i, row in enumerate(item["links"]):
            x = .65 + i * 4.12
            rect(slide, x, 5.0, 3.8, .75, ACID if i == 0 else PANEL, 0, None if i == 0 else LINE)
            add_text(slide, row[0], x + .14, 5.13, 3.5, .46, 15, BG if i == 0 else INK, True)
    elif kind == "source":
        title(slide, item["title"], 7.5, 1.45, 5.0, 1.1, 38)
        add_text(slide, item["body"], 7.65, 2.9, 4.8, 2.2, 18, MUTED)
        add_text(slide, item["note"], 7.65, 5.55, 4.8, .45, 12, ACID)
        add_text(slide, item["link"], .7, 6.15, 11.7, .35, 11, CYAN, False, PP_ALIGN.LEFT, mono=True)

prs.core_properties.title = DATA["title"] + " · " + DATA["subtitle"]
prs.core_properties.subject = "Three browser 3D games and their skill pack"
prs.core_properties.author = "Ariel Aizenshtat"
prs.core_properties.keywords = "Three.js, 3D games, browser game, Hebrew"
prs.save(OUT)
print(f"slides={len(prs.slides)} file={OUT} bytes={OUT.stat().st_size}")
